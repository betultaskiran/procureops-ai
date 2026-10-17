"""
app/services/extraction/excel_parser.py

Excel (.xlsx) dosyalarından yapılandırılmış metin ve tablo çıkarma servisi.

NEDEN OPENPYXL?
  - Modern .xlsx dosyaları için Python standardı
  - data_only=True ile formüllerin hesaplanmış sayısal değerlerini okur
  - read_only=True ile bellek tüketimini minimize eder
  - Birden fazla çalışma sayfasını (sheets) destekler
"""

from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as e:
    raise ImportError(
        "openpyxl is required for Excel parsing. "
        "Add it to pyproject.toml: 'openpyxl>=3.1.2'"
    ) from e


class ExcelParseError(Exception):
    """Excel dosyası açılamadı veya ayrıştırılamadı."""
    pass


class EmptyExcelError(ExcelParseError):
    """Excel çalışma kitabı boş veya okunabilir veri içermiyor."""
    pass


def _format_cell_value(val: Any) -> str:
    """Hücre değerini temiz string formatına dönüştürür."""
    if val is None:
        return ""
    if isinstance(val, float):
        # Gereksiz .0 float değerlerini temizle (ör. adet: 10.0 -> 10)
        if val.is_integer():
            return str(int(val))
        return f"{val:.4f}".rstrip("0").rstrip(".")
    return str(val).strip()


def extract_text_from_excel(file_path: str | Path) -> str:
    """
    Excel (.xlsx) dosyasındaki tüm çalışma sayfalarını okuyup
    LLM structured extraction için optimize edilmiş Markdown tablo formatına dönüştürür.

    Returns:
        --- Sheet: Sayfa1 ---
        Kolon1 | Kolon2 | Kolon3
        Değer1 | Değer2 | Değer3

    Raises:
        ExcelParseError: Dosya bozuk veya açılamıyor
        EmptyExcelError: Dosyada hiç veri yok
    """
    path = Path(file_path)
    if not path.exists():
        raise ExcelParseError(f"File not found: {path}")

    try:
        # data_only=True: Formül yerine hesaplanmış değeri al
        # read_only=True: Büyük dosyalarda hızlı ve az RAM
        wb = openpyxl.load_workbook(filename=str(path), data_only=True, read_only=True)
    except Exception as e:
        raise ExcelParseError(f"Failed to open Excel workbook: {e}") from e

    try:
        sheet_outputs: list[str] = []
        total_cells_with_data = 0

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows_data: list[list[str]] = []

            for row in sheet.iter_rows(values_only=True):
                if not row:
                    continue
                formatted_cells = [_format_cell_value(c) for c in row]
                # Satır tamamen boş değilse dahil et
                if any(formatted_cells):
                    rows_data.append(formatted_cells)
                    total_cells_with_data += sum(1 for c in formatted_cells if c)

            if rows_data:
                # Markdown tablosuna çevir
                # Tüm satırları en uzun satır boyutuna normalize et
                max_cols = max(len(r) for r in rows_data)
                table_lines = [f"--- Sheet: {sheet_name} ---"]
                for idx, r in enumerate(rows_data):
                    padded = r + [""] * (max_cols - len(r))
                    table_lines.append(" | ".join(padded))
                    if idx == 0 and len(rows_data) > 1:
                        # Başlık satırı altına ayırıcı
                        table_lines.append(" | ".join(["---"] * max_cols))

                sheet_outputs.append("\n".join(table_lines))

        wb.close()

        if total_cells_with_data == 0 or not sheet_outputs:
            raise EmptyExcelError("Excel workbook contains no readable data or is completely empty")

        return "\n\n".join(sheet_outputs)

    except EmptyExcelError:
        raise
    except Exception as e:
        raise ExcelParseError(f"Failed to parse Excel workbook contents: {e}") from e
